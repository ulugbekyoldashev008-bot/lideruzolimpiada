import os, json, asyncio, random, io, re, time, math
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo
from pathlib import Path

from dotenv import load_dotenv
from psycopg.rows import dict_row
from psycopg_pool import ConnectionPool
from telegram import Update, ReplyKeyboardMarkup, ReplyKeyboardRemove, KeyboardButton, InlineKeyboardMarkup, InlineKeyboardButton
from telegram.ext import Application, CommandHandler, MessageHandler, CallbackQueryHandler, ConversationHandler, ContextTypes, filters
from openpyxl import Workbook
from openpyxl.styles import Font

load_dotenv()
TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_LOGIN = os.getenv("ADMIN_LOGIN", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "").strip()
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
TZ_NAME = os.getenv("TIMEZONE", "Asia/Tashkent")
try:
    TZ = ZoneInfo(TZ_NAME)
except Exception:
    TZ = timezone(timedelta(hours=5))

SUBJECTS = ["Ingliz tili", "Matematika", "Mental arifmetika", "IT"]
# Ingliz tili darajalari (tests/questions.json dagi bo‘limlar)
ENGLISH_LEVELS = {"Starter": "Starter", "Beginner": "Beginner", "Elementary": "Elementary", "Intermediate": "Intermediate", "Advanced": "Advanced"}
MATH_LEVELS = [f"{n}-sinf" for n in range(1, 12)]
GRADES = [f"{n}-sinf" for n in range(1, 12)]
QUESTIONS_PER_TEST = 30
DEFAULT_DURATION_MIN = 60
IT_LEVELS = ["Frontend", "Backend Python"]
# Mental arifmetika: misollar ketma-ket chiqadi, javob raqam bilan yoziladi,
# test boshlangandan keyin 10 daqiqa beriladi (qaysi misolda bo‘lsa ham shunda yakunlanadi)
MENTAL = "Mental arifmetika"
MENTAL_MINUTES = 10

def load_questions(name):
    path = Path(__file__).parent / "tests" / name
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}

_eng = load_questions("questions.json")
_math = load_questions("matematika.json")  # {"1-sinf": [...], ..., "11-sinf": [...]}
_mental = load_questions("mental.json")  # {"daraja": [{"n", "q", "answer"}, ...]}
_it = load_questions("it.json")  # {"Frontend": [...], "Backend Python": [...]}
QBANK = {
    "Ingliz tili": {lvl: _eng.get(key, []) for lvl, key in ENGLISH_LEVELS.items()},
    "Matematika": {lvl: _math.get(lvl, []) for lvl in MATH_LEVELS},
    MENTAL: dict(_mental),
    "IT": {lvl: _it.get(lvl, []) for lvl in IT_LEVELS},
}
LEVELS = {subj: list(levels) for subj, levels in QBANK.items()}
QINDEX = {s: {lvl: {q["n"]: q for q in qs} for lvl, qs in lv.items()} for s, lv in QBANK.items()}

def school_label(v):
    if not v:
        return "maktab kiritilmagan"
    return f"{v}-maktab" if str(v).isdigit() else str(v)


REG_FIELDS = ["name", "birth", "school", "grade", "subject", "level"]
REG_LABELS = {
    "name": "Ism va familyangizni to‘liq kiriting:",
    "birth": "Tug‘ilgan sanangizni yil.oy.kun ko‘rinishida kiriting.\nMasalan: 2012.05.23",
    "school": "Qaysi maktabda o‘qiysiz?\nMaktab nomini yozing (masalan: 25-maktab yoki Navoiy nomidagi maktab):",
    "grade": "Nechinchi sinfda o‘qiysiz?",
    "subject": "Qaysi fandan test topshirmoqchisiz?",
    "level": "Qaysi darajadansiz? Tanlang:",
}
REGISTRATION, ADMIN_LOGIN_STATE, ADMIN_PASSWORD_STATE = range(3)
CANCEL = "❌ Bekor qilish"

# ───────────────────────── Ma’lumotlar bazasi (PostgreSQL) ─────────────────────────
POOL = None

def db():
    return POOL.connection()

def q_one(sql, params=()):
    with db() as c:
        return c.execute(sql, params).fetchone()

def q_all(sql, params=()):
    with db() as c:
        return c.execute(sql, params).fetchall()

def q_exec(sql, params=()):
    with db() as c:
        return c.execute(sql, params).rowcount

SCHEMA = [
    """CREATE TABLE IF NOT EXISTS users(
        telegram_id BIGINT PRIMARY KEY, username TEXT, full_name TEXT NOT NULL, birth_date TEXT, age INTEGER,
        grade TEXT, school TEXT, subject TEXT, level TEXT, registered_at TEXT)""",
    # hozir ishlanayotgan testlar (qayta ishga tushganda ham yo‘qolmaydi)
    """CREATE TABLE IF NOT EXISTS exams(
        telegram_id BIGINT PRIMARY KEY, subject TEXT, level TEXT, question_ids TEXT,
        idx INTEGER NOT NULL DEFAULT 0, correct INTEGER NOT NULL DEFAULT 0,
        started_ts DOUBLE PRECISION, deadline_ts DOUBLE PRECISION)""",
    # yakunlangan testlar
    """CREATE TABLE IF NOT EXISTS results(
        id BIGSERIAL PRIMARY KEY, telegram_id BIGINT UNIQUE, full_name TEXT, subject TEXT, level TEXT,
        correct INTEGER, answered INTEGER, total INTEGER, percent DOUBLE PRECISION,
        started_at TEXT, finished_at TEXT, duration_sec INTEGER, reason TEXT)""",
    """CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT)""",
    """CREATE TABLE IF NOT EXISTS admins(telegram_id BIGINT PRIMARY KEY)""",
]

def init_db():
    global POOL
    POOL = ConnectionPool(DATABASE_URL, min_size=1, max_size=5, kwargs={"row_factory": dict_row}, open=True)
    with db() as c:
        for stmt in SCHEMA:
            c.execute(stmt)
        c.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS school TEXT")
        # "Pre-Intermediate" darajasi "Intermediate" ga almashtirildi
        c.execute("UPDATE users SET level='Intermediate' WHERE subject='Ingliz tili' AND level='Pre-Intermediate'")
        # mental arifmetika misollari almashtirildi: eski daraja nomi yangisiga o‘tkaziladi
        if _mental:
            c.execute("UPDATE users SET level=%s WHERE subject=%s AND level<>ALL(%s)", (list(_mental)[0], MENTAL, list(_mental)))
        for k, v in {"registration_enabled": "1", "start_at": "", "duration_min": str(DEFAULT_DURATION_MIN), "notified_for": ""}.items():
            c.execute("INSERT INTO settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO NOTHING", (k, v))

def get_setting(key, default=""):
    row = q_one("SELECT value FROM settings WHERE key=%s", (key,))
    return row["value"] if row else default

def set_setting(key, value):
    q_exec("INSERT INTO settings(key,value) VALUES(%s,%s) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value", (key, str(value)))

def get_user(uid):
    return q_one("SELECT * FROM users WHERE telegram_id=%s", (uid,))

def get_exam(uid):
    return q_one("SELECT * FROM exams WHERE telegram_id=%s", (uid,))

def has_result(uid):
    return q_one("SELECT 1 AS x FROM results WHERE telegram_id=%s", (uid,)) is not None

def is_admin(uid):
    return q_one("SELECT 1 AS x FROM admins WHERE telegram_id=%s", (uid,)) is not None

# ───────────────────────── Vaqt yordamchilari ─────────────────────────
def test_time():
    val = get_setting("start_at", "").strip()
    if not val:
        return None
    try:
        return datetime.strptime(val, "%Y-%m-%d %H:%M").replace(tzinfo=TZ)
    except ValueError:
        return None

def duration_min():
    try:
        return max(1, int(get_setting("duration_min", str(DEFAULT_DURATION_MIN))))
    except ValueError:
        return DEFAULT_DURATION_MIN

def test_end():
    start = test_time()
    return start + timedelta(minutes=duration_min()) if start else None

def fmt_dur(sec):
    m, s = divmod(int(sec or 0), 60)
    return f"{m} daq {s:02d} s"

def exam_deadline(subject, started_ts, end_ts):
    """Mental arifmetikada boshlangandan MENTAL_MINUTES daqiqa (olimpiada tugashidan oshmaydi), boshqalarda olimpiada tugashi."""
    return min(end_ts, started_ts + MENTAL_MINUTES * 60) if subject == MENTAL else end_ts

def minutes_left(deadline_ts):
    return max(0, math.ceil((deadline_ts - time.time()) / 60))

def stamp(ts=None):
    return datetime.fromtimestamp(ts if ts is not None else time.time(), TZ).strftime("%d.%m.%Y %H:%M:%S")

# ───────────────────────── Klaviaturalar ─────────────────────────
def main_keyboard():
    return ReplyKeyboardMarkup([["📝 Ro‘yxatdan o‘tish"], ["🔐 Admin panelga kirish"]], resize_keyboard=True)

def home_keyboard(uid=None):
    if uid is not None and has_result(uid):
        return ReplyKeyboardMarkup([["👤 Profil"], ["🔐 Admin panelga kirish"]], resize_keyboard=True)
    return ReplyKeyboardMarkup([["🧪 Testni boshlash"], ["👤 Profil"], ["🔐 Admin panelga kirish"]], resize_keyboard=True)

def admin_keyboard():
    enabled = get_setting("registration_enabled", "1") == "1"
    return ReplyKeyboardMarkup([
        ["📊 Statistika", "📥 Excel yuklab olish"],
        ["✅ Topshirganlar", "🔄 Ishlayotganlar"],
        ["⏳ Topshirmaganlar"],
        ["🟢 Ro‘yxatdan o‘tishni " + ("o‘chirish" if enabled else "yoqish")],
        ["🕒 Olimpiada vaqtini belgilash", "⏱ Davomiylikni o‘zgartirish"],
        ["📅 Olimpiada vaqtini ko‘rish"],
        ["🚪 Admin paneldan chiqish"],
    ], resize_keyboard=True)

def chunk(items, n=2):
    return [items[i:i + n] for i in range(0, len(items), n)]

def reg_keyboard(key, reg=None):
    reg = reg or {}
    rows = []
    if key == "grade":
        rows = chunk(GRADES, 3)
    elif key == "subject":
        rows = [[s_] for s_ in SUBJECTS]
    elif key == "level":
        rows = chunk(LEVELS.get(reg.get("subject"), []), 2 if reg.get("subject") == "Ingliz tili" else 3)
    return ReplyKeyboardMarkup(rows + [[CANCEL]], resize_keyboard=True)

# ───────────────────────── Ro‘yxatdan o‘tish va kabinet ─────────────────────────
def parse_birth(text):
    try:
        d = datetime.strptime(re.sub(r"[-/ ]", ".", text.strip()), "%Y.%m.%d").date()
    except ValueError:
        return None
    today = datetime.now(TZ).date()
    age = today.year - d.year - ((today.month, today.day) < (d.month, d.day))
    return d if 4 <= age <= 100 else None

def age_from(birth_iso):
    d = datetime.strptime(birth_iso, "%Y-%m-%d").date()
    today = datetime.now(TZ).date()
    return today.year - d.year - ((today.month, today.day) < (d.month, d.day))

def schedule_line():
    start, end = test_time(), test_end()
    if not start:
        return "Olimpiada vaqti hali belgilanmagan"
    return f"{start.strftime('%d.%m.%Y')}, soat {start.strftime('%H:%M')} – {end.strftime('%H:%M')} (Toshkent vaqti, {duration_min()} daqiqa)"

def profile_text(u, title="👤 Shaxsiy kabinet"):
    birth = u["birth_date"].replace("-", ".") if u["birth_date"] else "—"
    if has_result(u["telegram_id"]):
        status = "✅ Test topshirilgan"
    elif get_exam(u["telegram_id"]):
        status = "🔄 Test jarayonida"
    else:
        status = f"⏳ Olimpiada: {schedule_line()}"
    return (f"{title}\n\nIsm-familya: {u['full_name']}\nTug‘ilgan sana: {birth} ({u['age']} yosh)\n"
            f"Maktab: {school_label(u['school'])}\nSinf: {u['grade']}\n"
            f"Fan: {u['subject'] or '—'}\nDaraja: {u['level'] or '—'}\n\n{status}\n"
            "Test vaqti kelganda tanlagan fan va darajangiz bo‘yicha test beriladi.")

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    u = get_user(update.effective_user.id)
    if u:
        await update.message.reply_text(profile_text(u), reply_markup=home_keyboard(u["telegram_id"]))
    else:
        await update.message.reply_text("Assalomu alaykum! Olimpiada botiga xush kelibsiz.\nQuyidagi tugmalardan birini tanlang.", reply_markup=main_keyboard())
    return ConversationHandler.END

async def registration_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    if has_result(uid) or get_exam(uid):
        await update.message.reply_text("Siz testni boshlagansiz yoki topshirib bo‘lgansiz, qayta ro‘yxatdan o‘tish shart emas.", reply_markup=home_keyboard(uid))
        return ConversationHandler.END
    if get_setting("registration_enabled", "1") != "1":
        await update.message.reply_text("Hozir ro‘yxatdan o‘tish yopilgan. Admin ochganda qayta urinib ko‘ring.", reply_markup=main_keyboard())
        return ConversationHandler.END
    context.user_data["reg"] = {}
    context.user_data["reg_index"] = 0
    await update.message.reply_text(REG_LABELS["name"], reply_markup=reg_keyboard("name"))
    return REGISTRATION

async def registration_step(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    if text == CANCEL:
        await update.message.reply_text("Ro‘yxatdan o‘tish bekor qilindi.", reply_markup=main_keyboard())
        return ConversationHandler.END
    reg = context.user_data.setdefault("reg", {})
    i = context.user_data.get("reg_index", 0)
    key = REG_FIELDS[i]
    error = None
    if key == "name" and len(text.split()) < 2:
        error = "Iltimos, ism va familyani birga kiriting."
    elif key == "birth":
        d = parse_birth(text)
        if d:
            text = d.strftime("%Y-%m-%d")
        else:
            error = "Sana noto‘g‘ri. Yil.oy.kun ko‘rinishida kiriting, masalan: 2012.05.23"
    elif key == "school":
        text = " ".join(text.split())
        if len(text) < 2 or len(text) > 100:
            error = "Maktab nomini to‘liq yozing (kamida 2 ta belgi), masalan: 25-maktab"
    elif key == "grade" and text not in GRADES:
        error = "Iltimos, sinfni tugmalardan tanlang."
    elif key == "subject" and text not in SUBJECTS:
        error = "Iltimos, fanni tugmalardan tanlang."
    elif key == "level" and text not in LEVELS.get(reg.get("subject"), []):
        error = "Iltimos, darajani tugmalardan tanlang."
    if error:
        await update.message.reply_text(error)
        return REGISTRATION
    reg[key] = text
    i += 1
    context.user_data["reg_index"] = i
    if i < len(REG_FIELDS):
        await update.message.reply_text(REG_LABELS[REG_FIELDS[i]], reply_markup=reg_keyboard(REG_FIELDS[i], reg))
        return REGISTRATION
    user = update.effective_user
    q_exec("""INSERT INTO users(telegram_id,username,full_name,birth_date,age,grade,school,subject,level,registered_at)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(telegram_id) DO UPDATE SET
        username=EXCLUDED.username, full_name=EXCLUDED.full_name, birth_date=EXCLUDED.birth_date, age=EXCLUDED.age,
        grade=EXCLUDED.grade, school=EXCLUDED.school, subject=EXCLUDED.subject,
        level=EXCLUDED.level, registered_at=EXCLUDED.registered_at""",
        (user.id, user.username, reg["name"], reg["birth"], age_from(reg["birth"]), reg["grade"],
         reg["school"], reg["subject"], reg["level"], stamp()))
    context.user_data.pop("reg", None)
    context.user_data.pop("reg_index", None)
    await update.message.reply_text(profile_text(get_user(user.id), "✅ Ro‘yxatdan o‘tish yakunlandi!\n\n👤 Shaxsiy kabinetingiz ochildi"),
                                    reply_markup=home_keyboard(user.id))
    return ConversationHandler.END

async def profile(update, context):
    u = get_user(update.effective_user.id)
    if not u:
        await update.message.reply_text("Avval ro‘yxatdan o‘ting.", reply_markup=main_keyboard()); return
    await update.message.reply_text(profile_text(u), reply_markup=home_keyboard(u["telegram_id"]))

# ───────────────────────── Test ─────────────────────────
async def start_exam(uid, context):
    """O‘quvchi ro‘yxatdan o‘tishda tanlagan fan va daraja bo‘yicha testni boshlaydi."""
    bot = context.bot
    send = lambda text, **kw: bot.send_message(uid, text, **kw)
    u = get_user(uid)
    if not u:
        await send("Avval ro‘yxatdan o‘ting.", reply_markup=main_keyboard()); return
    if has_result(uid):
        await send("✅ Siz olimpiada testini allaqachon topshirgansiz. Qayta topshirish mumkin emas.", reply_markup=home_keyboard(uid)); return
    start, end = test_time(), test_end()
    if not start:
        await send("⏳ Olimpiada boshlanish vaqti hali e’lon qilinmagan. Test admin belgilagan vaqtdan keyin boshlanadi."); return
    now = time.time()
    if now < start.timestamp():
        await send(f"⏳ Olimpiada hali boshlanmadi.\n\n📅 Sana: {start.strftime('%d.%m.%Y')}\n🕒 Vaqt: {start.strftime('%H:%M')} (Toshkent vaqti)\n"
                   f"⌛ Davomiyligi: {duration_min()} daqiqa\n\nTest aynan shu sana va vaqtda boshlanadi. Shu vaqtda “🧪 Testni boshlash” tugmasini bosing "
                   "yoki botdan kelgan xabardagi tugmadan foydalaning."); return
    if now >= end.timestamp():
        await send("⌛ Olimpiada yakunlangan. Test ishlash vaqti tugadi.", reply_markup=home_keyboard(uid)); return
    if get_exam(uid):
        await send("Test davom etmoqda. Joriy savolga javob bering.")
        await send_question(uid, bot); return
    subject, level = u["subject"], u["level"]
    if subject not in QINDEX or level not in QINDEX[subject]:
        await send("Ro‘yxatdan o‘tish ma’lumotlaringiz eskirgan. Iltimos, qaytadan ro‘yxatdan o‘ting.", reply_markup=main_keyboard()); return
    pool = list(QINDEX[subject][level])
    if not pool:
        await send(f"“{subject} — {level}” uchun savollar hali qo‘shilmagan. Iltimos, admin bilan bog‘laning."); return
    ids = pool if subject == MENTAL else random.sample(pool, min(QUESTIONS_PER_TEST, len(pool)))
    deadline = exam_deadline(subject, now, end.timestamp())
    created = q_one("""INSERT INTO exams(telegram_id,subject,level,question_ids,idx,correct,started_ts,deadline_ts)
        VALUES(%s,%s,%s,%s,0,0,%s,%s) ON CONFLICT(telegram_id) DO NOTHING RETURNING telegram_id""",
        (uid, subject, level, json.dumps(ids), now, deadline))
    if not created:
        await send_question(uid, bot); return
    if subject == MENTAL:
        asyncio.create_task(expire_at(bot, uid, deadline))
    how = "Javobni raqam bilan yozib yuboring." if subject == MENTAL else "Javobni A/B/C/D tugmalaridan tanlang."
    await send(f"📚 {subject} — {level}\nSavollar soni: {len(ids)}\n"
               f"⏱ Test {datetime.fromtimestamp(deadline, TZ).strftime('%H:%M')} gacha davom etadi (qolgan vaqt: {minutes_left(deadline)} daqiqa).\n"
               f"{how}\n\nOmad! Boshladik!", reply_markup=ReplyKeyboardRemove())
    await send_question(uid, bot)

async def test_start(update, context):
    await start_exam(update.effective_user.id, context)

async def begin_callback(update, context):
    query = update.callback_query
    await query.answer()
    await start_exam(query.from_user.id, context)

async def send_question(uid, bot):
    exam = get_exam(uid)
    if not exam:
        return
    ids = json.loads(exam["question_ids"])
    i = exam["idx"]
    q = QINDEX.get(exam["subject"], {}).get(exam["level"], {}).get(ids[i]) if i < len(ids) else None
    if q is None:
        await finish_exam(bot, uid); return
    text = q["q"]
    if "answer" in q:  # mental arifmetika: javob yoziladi, tugma yo‘q
        await bot.send_message(uid, f"❓ {i+1}/{len(ids)}  ⏱ qolgan vaqt: ~{minutes_left(exam['deadline_ts'])} daq\n\n{text}\n\n✍️ Javobni raqam bilan yozing.")
        return
    if any(len(option) > 40 for option in q["options"]):
        # uzun variantlar tugmaga sig‘maydi: matnda ko‘rsatiladi, tugmada faqat harf
        text += "\n\n" + "\n".join(f"{'ABCD'[k]}) {option}" for k, option in enumerate(q["options"]))
        opts = [[InlineKeyboardButton('ABCD'[k], callback_data=f"answer:{i}:{k}") for k in range(len(q["options"]))]]
    else:
        opts = [[InlineKeyboardButton(f"{'ABCD'[k]}. {option}", callback_data=f"answer:{i}:{k}")] for k, option in enumerate(q["options"])]
    await bot.send_message(uid, f"❓ {i+1}/{len(ids)}  ⏱ qolgan vaqt: ~{minutes_left(exam['deadline_ts'])} daq\n\n{text}",
                           reply_markup=InlineKeyboardMarkup(opts))

async def answer_question(update, context):
    query = update.callback_query
    await query.answer()
    uid = query.from_user.id
    exam = get_exam(uid)
    if not exam:
        await query.edit_message_text("Test topilmadi yoki yakunlangan."); return
    if time.time() >= exam["deadline_ts"]:
        await query.edit_message_text("⏰ Vaqt tugadi.")
        await finish_exam(context.bot, uid, expired=True); return
    _, qi, ai = query.data.split(":")
    qi, ai = int(qi), int(ai)
    if qi != exam["idx"]:
        return  # eski yoki ikki marta bosilgan tugma
    ids = json.loads(exam["question_ids"])
    q = QINDEX[exam["subject"]][exam["level"]][ids[qi]]
    ok = 1 if "ABCD"[ai] == q["correct"] else 0
    changed = q_exec("UPDATE exams SET idx=idx+1, correct=correct+%s WHERE telegram_id=%s AND idx=%s", (ok, uid, qi))
    if not changed:
        return
    await query.edit_message_text("✅ Javob qabul qilindi." if ok else "Javob qabul qilindi.")
    await send_question(uid, context.bot)

async def answer_typed(update, context, exam):
    """Mental arifmetika: o‘quvchi javobni raqam bilan yozadi."""
    uid = update.effective_user.id
    text = (update.message.text or "").strip().replace(" ", "")
    if time.time() >= exam["deadline_ts"]:
        await update.message.reply_text("⏰ Vaqt tugadi.")
        await finish_exam(context.bot, uid, expired=True); return
    if not re.fullmatch(r"-?\d+", text):
        await update.message.reply_text("Javobni faqat raqam bilan yozing (masalan: 364)."); return
    ids = json.loads(exam["question_ids"])
    qi = exam["idx"]
    q = QINDEX[exam["subject"]][exam["level"]][ids[qi]]
    ok = 1 if int(text) == q["answer"] else 0
    if not q_exec("UPDATE exams SET idx=idx+1, correct=correct+%s WHERE telegram_id=%s AND idx=%s", (ok, uid, qi)):
        return
    await send_question(uid, context.bot)

async def finish_exam(bot, uid, expired=False):
    # DELETE ... RETURNING testni atomik “egallaydi”: ikki marta yakunlanib ketmaydi
    exam = q_one("DELETE FROM exams WHERE telegram_id=%s RETURNING *", (uid,))
    if not exam:
        return
    u = get_user(uid)
    name = u["full_name"] if u else str(uid)
    total = len(json.loads(exam["question_ids"]))
    correct, answered = exam["correct"], exam["idx"]
    finished_ts = exam["deadline_ts"] if expired else min(time.time(), exam["deadline_ts"])
    duration = max(0, int(finished_ts - exam["started_ts"]))
    mental = exam["subject"] == MENTAL
    # mental arifmetikada foiz yo‘q: har bir to‘g‘ri javob 1 ball
    percent = None if mental else (round(correct * 100 / total, 2) if total else 0)
    reason = "Vaqt tugadi" if expired else "Tugatdi"
    q_exec("""INSERT INTO results(telegram_id,full_name,subject,level,correct,answered,total,percent,started_at,finished_at,duration_sec,reason)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT(telegram_id) DO NOTHING""",
        (uid, name, exam["subject"], exam["level"], correct, answered, total, percent,
         stamp(exam["started_ts"]), stamp(finished_ts), duration, reason))
    head = "⏰ Vaqt tugadi, test yakunlandi!" if expired else "🏁 Test yakunlandi!"
    unanswered = total - answered
    extra = f"\nJavob berilmagan: {unanswered} ta (noto‘g‘ri hisoblandi)" if unanswered else ""
    if mental:
        body = (f"Ishlangan misollar: {answered} ta\nTo‘g‘ri javoblar: {correct} ta\nNoto‘g‘ri javoblar: {answered-correct} ta\n"
                f"🏅 Ball: {correct}")
    else:
        body = (f"To‘g‘ri javoblar: {correct}/{total}\nNoto‘g‘ri javoblar: {total-correct}/{total}{extra}\n"
                f"Natija: {percent}%")
    msg = await bot.send_message(uid, f"{head}\n\nO‘quvchi: {name}\nFan: {exam['subject']}\nDaraja: {exam['level']}\n"
                                      f"{body}\nSarflangan vaqt: {fmt_dur(duration)}", reply_markup=home_keyboard(uid))
    async def delete_later():
        await asyncio.sleep(15)
        try:
            await bot.delete_message(chat_id=uid, message_id=msg.message_id)
        except Exception as e:
            print("Natija xabarini o‘chirishda xatolik:", repr(e))
    asyncio.create_task(delete_later())

# ───────────────────────── Fon jarayoni: boshlanish xabari va vaqt tugashi ─────────────────────────
async def expire_exams(bot):
    for r in q_all("SELECT telegram_id FROM exams WHERE deadline_ts <= %s", (time.time(),)):
        try:
            await finish_exam(bot, r["telegram_id"], expired=True)
        except Exception as e:
            print("Testni yakunlashda xatolik:", r["telegram_id"], repr(e))

async def expire_at(bot, uid, deadline_ts):
    """Mental arifmetika testini aynan vaqti tugaganda yakunlaydi (watcher esa zaxira sifatida qoladi)."""
    await asyncio.sleep(max(0, deadline_ts - time.time()))
    exam = get_exam(uid)
    if exam and exam["deadline_ts"] <= time.time():
        try:
            await finish_exam(bot, uid, expired=True)
        except Exception as e:
            print("Testni yakunlashda xatolik:", uid, repr(e))

async def notify_start(bot):
    start, end = test_time(), test_end()
    val = get_setting("start_at", "")
    if not start or time.time() < start.timestamp() or get_setting("notified_for", "") == val:
        return
    set_setting("notified_for", val)
    if time.time() >= end.timestamp():
        return
    rows = q_all("""SELECT u.telegram_id, u.subject, u.level FROM users u
        WHERE NOT EXISTS (SELECT 1 FROM results r WHERE r.telegram_id=u.telegram_id)
          AND NOT EXISTS (SELECT 1 FROM exams e WHERE e.telegram_id=u.telegram_id)""")
    for r in rows:
        try:
            await bot.send_message(r["telegram_id"],
                f"🔔 Olimpiada boshlandi!\nFan: {r['subject'] or '—'}\nDaraja: {r['level'] or '—'}\n"
                + (f"Testni {end.strftime('%H:%M')} gacha boshlashingiz mumkin. Boshlaganingizdan keyin {MENTAL_MINUTES} daqiqa vaqt beriladi."
                 if r["subject"] == MENTAL else f"Test {end.strftime('%H:%M')} gacha davom etadi ({duration_min()} daqiqa).")
                + " Tayyor bo‘lsangiz, tugmani bosing.",
                reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("🧪 Testni boshlash", callback_data="begin:go")]]))
        except Exception as e:
            print("Xabar yuborishda xatolik:", r["telegram_id"], repr(e))
        await asyncio.sleep(0.05)

async def watcher(app):
    while True:
        await asyncio.sleep(15)
        try:
            await expire_exams(app.bot)
            await notify_start(app.bot)
        except Exception as e:
            print("watcher xatosi:", repr(e))

# ───────────────────────── Admin ─────────────────────────
FAILS = {}  # uid -> (urinishlar soni, bloklangan vaqt)

async def admin_begin(update, context):
    fails, locked = FAILS.get(update.effective_user.id, (0, 0))
    if time.time() < locked:
        await update.message.reply_text("Juda ko‘p noto‘g‘ri urinish. 10 daqiqadan keyin qayta urinib ko‘ring.", reply_markup=main_keyboard())
        return ConversationHandler.END
    await update.message.reply_text("Admin loginini kiriting:", reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    return ADMIN_LOGIN_STATE

async def admin_login(update, context):
    if update.message.text == CANCEL:
        await update.message.reply_text("Bekor qilindi.", reply_markup=main_keyboard()); return ConversationHandler.END
    context.user_data["admin_login"] = update.message.text.strip()
    await update.message.reply_text("Admin parolini kiriting:")
    return ADMIN_PASSWORD_STATE

async def admin_password(update, context):
    uid = update.effective_user.id
    if update.message.text == CANCEL:
        await update.message.reply_text("Bekor qilindi.", reply_markup=main_keyboard()); return ConversationHandler.END
    if context.user_data.get("admin_login") != ADMIN_LOGIN or update.message.text.strip() != ADMIN_PASSWORD:
        fails = FAILS.get(uid, (0, 0))[0] + 1
        FAILS[uid] = (fails, time.time() + 600 if fails >= 5 else 0)
        if fails >= 5:
            FAILS[uid] = (0, time.time() + 600)
        await update.message.reply_text("❌ Login yoki parol noto‘g‘ri.", reply_markup=main_keyboard())
        return ConversationHandler.END
    FAILS.pop(uid, None)
    q_exec("INSERT INTO admins(telegram_id) VALUES(%s) ON CONFLICT DO NOTHING", (uid,))
    await update.message.reply_text("✅ Admin panelga xush kelibsiz.", reply_markup=admin_keyboard())
    return ConversationHandler.END

async def admin_exit(update, context):
    if not is_admin(update.effective_user.id): return
    q_exec("DELETE FROM admins WHERE telegram_id=%s", (update.effective_user.id,))
    await update.message.reply_text("Admin paneldan chiqdingiz.", reply_markup=main_keyboard())

def olympiad_state():
    start, end = test_time(), test_end()
    if not start: return "Vaqt belgilanmagan"
    now = time.time()
    if now < start.timestamp(): return "Hali boshlanmagan"
    if now < end.timestamp(): return "Davom etmoqda"
    return "Yakunlangan"

async def admin_stats(update, context):
    if not is_admin(update.effective_user.id): return
    total = q_one("SELECT COUNT(*) AS n FROM users")["n"]
    done = q_one("SELECT COUNT(*) AS n FROM results")["n"]
    active = q_one("SELECT COUNT(*) AS n FROM exams")["n"]
    avg = q_one("SELECT AVG(percent) AS a FROM results WHERE percent IS NOT NULL")["a"]
    avg_ball = q_one("SELECT AVG(correct) AS a FROM results WHERE subject=%s", (MENTAL,))["a"]
    waiting = max(0, total - done - active)
    reg_by = {(r["subject"], r["level"]): r["n"] for r in q_all("SELECT subject, level, COUNT(*) AS n FROM users GROUP BY subject, level")}
    done_by = {(r["subject"], r["level"]): r["n"] for r in q_all("SELECT subject, level, COUNT(*) AS n FROM results GROUP BY subject, level")}
    lines = []
    for subj in SUBJECTS:
        for lvl in LEVELS[subj]:
            n = reg_by.get((subj, lvl), 0)
            if n:
                lines.append(f"  • {subj} — {lvl}: {n} ta ro‘yxatdan, {done_by.get((subj, lvl), 0)} ta topshirgan")
    await update.message.reply_text(
        f"📊 Statistika\n\nHolat: {olympiad_state()}\nOlimpiada: {schedule_line()}\n\n"
        f"👥 Ro‘yxatdan o‘tganlar: {total}\n✅ Test topshirganlar: {done}\n🔄 Hozir ishlayotganlar: {active}\n⏳ Hali topshirmaganlar: {waiting}\n"
        + (f"📈 O‘rtacha natija: {float(avg):.1f}%\n" if avg is not None else "")
        + (f"🧮 {MENTAL}, o‘rtacha ball: {float(avg_ball):.1f}\n" if avg_ball is not None else "")
        + ("\nFan va daraja bo‘yicha:\n" + "\n".join(lines) if lines else ""),
        reply_markup=admin_keyboard())

def score_text(r):
    if r["subject"] == MENTAL:
        return f"{r['correct']} ball ({r['answered']} ta misol ishlagan)"
    return f"{r['percent']}% ({r['correct']}/{r['total']})"

async def send_lines(update, header, lines):
    await update.message.reply_text(header, reply_markup=admin_keyboard())
    if not lines:
        await update.message.reply_text("Hech kim yo‘q."); return
    buf = ""
    for line in lines:
        if len(buf) + len(line) + 1 > 3500:
            await update.message.reply_text(buf); buf = ""
        buf += line + "\n"
    if buf:
        await update.message.reply_text(buf)

async def admin_finished(update, context):
    if not is_admin(update.effective_user.id): return
    rows = q_all("SELECT * FROM results ORDER BY subject, level, percent DESC NULLS LAST, correct DESC, duration_sec ASC")
    lines = [f"{i}. {r['full_name']} — {r['subject']}/{r['level']}: {score_text(r)}, {fmt_dur(r['duration_sec'])}"
             for i, r in enumerate(rows, 1)]
    await send_lines(update, f"✅ Test topshirganlar: {len(rows)} ta", lines)

async def admin_active(update, context):
    if not is_admin(update.effective_user.id): return
    rows = q_all("""SELECT e.*, u.full_name FROM exams e LEFT JOIN users u ON u.telegram_id=e.telegram_id ORDER BY e.started_ts""")
    lines = [f"{i}. {r['full_name']} — {r['subject']}/{r['level']}: {r['idx']}/{len(json.loads(r['question_ids']))} savol, qolgan ~{minutes_left(r['deadline_ts'])} daq"
             for i, r in enumerate(rows, 1)]
    await send_lines(update, f"🔄 Hozir ishlayotganlar: {len(rows)} ta", lines)

WAITING_SQL = """SELECT u.* FROM users u
    WHERE NOT EXISTS (SELECT 1 FROM results r WHERE r.telegram_id=u.telegram_id)
      AND NOT EXISTS (SELECT 1 FROM exams e WHERE e.telegram_id=u.telegram_id)
    ORDER BY LOWER(u.full_name)"""

async def admin_waiting(update, context):
    if not is_admin(update.effective_user.id): return
    rows = q_all(WAITING_SQL)
    lines = [f"{i}. {r['full_name']} — {r['subject']}/{r['level']} ({school_label(r['school'])})" for i, r in enumerate(rows, 1)]
    await send_lines(update, f"⏳ Hali test ishlamaganlar: {len(rows)} ta", lines)

async def toggle_registration(update, context):
    if not is_admin(update.effective_user.id): return
    val = "0" if get_setting("registration_enabled", "1") == "1" else "1"
    set_setting("registration_enabled", val)
    await update.message.reply_text("Ro‘yxatdan o‘tish " + ("yoqildi." if val == "1" else "o‘chirildi."), reply_markup=admin_keyboard())

async def schedule_prompt(update, context):
    if not is_admin(update.effective_user.id): return
    await update.message.reply_text("Olimpiada boshlanish vaqtini kiriting: YYYY-MM-DD HH:MM\nMasalan: 2026-10-20 17:00\nVaqt Asia/Tashkent bo‘yicha.",
                                    reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    context.user_data["awaiting_schedule"] = "start"

async def duration_prompt(update, context):
    if not is_admin(update.effective_user.id): return
    await update.message.reply_text(f"Test necha daqiqa davom etadi? Hozir: {duration_min()} daqiqa.\nYangi qiymatni raqam bilan kiriting (masalan: 60).",
                                    reply_markup=ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True))
    context.user_data["awaiting_schedule"] = "duration_only"

def apply_new_deadline():
    """Ishlayotgan testlar muddatini yangi tugash vaqtiga moslaydi."""
    end = test_end()
    if end:
        q_exec("UPDATE exams SET deadline_ts=%s WHERE subject<>%s", (end.timestamp(), MENTAL))
        q_exec("UPDATE exams SET deadline_ts=LEAST(started_ts+%s, %s) WHERE subject=%s", (MENTAL_MINUTES * 60, end.timestamp(), MENTAL))

async def schedule_save(update, context):
    state = context.user_data.get("awaiting_schedule")
    if not state: return
    text = (update.message.text or "").strip()
    cancel_kb = ReplyKeyboardMarkup([[CANCEL]], resize_keyboard=True)
    if text == CANCEL:
        context.user_data.pop("awaiting_schedule", None)
        await update.message.reply_text("Bekor qilindi.", reply_markup=admin_keyboard()); return
    if state == "start":
        try:
            dt = datetime.strptime(text, "%Y-%m-%d %H:%M")
        except ValueError:
            await update.message.reply_text("Format noto‘g‘ri. YYYY-MM-DD HH:MM ko‘rinishida kiriting.", reply_markup=cancel_kb); return
        if dt < datetime.now(TZ).replace(tzinfo=None):
            await update.message.reply_text("O‘tgan vaqtni belgilab bo‘lmaydi. Qayta kiriting.", reply_markup=cancel_kb); return
        context.user_data["new_start"] = text
        context.user_data["awaiting_schedule"] = "duration"
        await update.message.reply_text("Test necha daqiqa davom etadi? Raqam bilan kiriting (masalan: 60).\nShu vaqt tugagach test to‘xtatiladi.", reply_markup=cancel_kb)
        return
    if not text.isdigit() or not 1 <= int(text) <= 1440:
        await update.message.reply_text("Davomiylikni daqiqalarda, 1 dan 1440 gacha raqam bilan kiriting.", reply_markup=cancel_kb); return
    if state == "duration":
        set_setting("start_at", context.user_data.pop("new_start"))
    set_setting("duration_min", int(text))
    apply_new_deadline()
    context.user_data.pop("awaiting_schedule", None)
    await update.message.reply_text(f"✅ Saqlandi.\nOlimpiada: {schedule_line()}", reply_markup=admin_keyboard())

async def schedule_show(update, context):
    if not is_admin(update.effective_user.id): return
    await update.message.reply_text(f"Holat: {olympiad_state()}\nOlimpiada: {schedule_line()}", reply_markup=admin_keyboard())

def write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max((len(str(c.value if c.value is not None else "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 34)

def level_order(subject, level):
    try:
        return LEVELS[subject].index(level)
    except (KeyError, ValueError):
        return 99

async def export_excel(update, context):
    if not is_admin(update.effective_user.id): return
    users = q_all("SELECT * FROM users ORDER BY LOWER(full_name)")
    results = q_all("SELECT * FROM results")
    exams = q_all("SELECT * FROM exams")
    udict = {u["telegram_id"]: u for u in users}
    rdict = {r["telegram_id"]: r for r in results}
    edict = {e["telegram_id"]: e for e in exams}
    wb = Workbook()

    # 1) Natijalar: har bir fan/daraja bo‘yicha alohida o‘rin (foiz kamayishi, keyin kam vaqt)
    ws = wb.active; ws.title = "Natijalar"
    ordered = sorted(results, key=lambda r: (SUBJECTS.index(r["subject"]) if r["subject"] in SUBJECTS else 99,
                                              level_order(r["subject"], r["level"]), -(r["percent"] or 0), -(r["correct"] or 0), r["duration_sec"] or 0))
    rows, place, prev = [], 0, None
    for r in ordered:
        grp = (r["subject"], r["level"])
        place = place + 1 if grp == prev else 1
        prev = grp
        u = udict.get(r["telegram_id"]) or {}
        rows.append([place, r["full_name"], u.get("birth_date"), u.get("grade"), u.get("school"),
                     r["subject"], r["level"], r["correct"], (r["answered"] if r["subject"] == MENTAL else r["total"]) - r["correct"],
                     r["answered"], r["total"], r["percent"], r["correct"] if r["subject"] == MENTAL else None,
                     fmt_dur(r["duration_sec"]), r["duration_sec"], r["started_at"], r["finished_at"], r["reason"],
                     u.get("username"), r["telegram_id"]])
    write_sheet(ws, ["O‘rin (fan/daraja ichida)", "Ism-familya", "Tug‘ilgan sana", "Sinf", "Maktab", "Fan", "Daraja",
                     "To‘g‘ri", "Noto‘g‘ri", "Javob berilgan", "Jami savol", "Foiz (%)", "Ball (Mental arifmetika)", "Sarflangan vaqt", "Sarflangan vaqt (soniya)",
                     "Boshlangan", "Tugagan", "Tugash sababi", "Username", "Telegram ID"], rows)

    # 2) Barcha ro‘yxatdan o‘tganlar va holati
    ws2 = wb.create_sheet("Barcha o‘quvchilar")
    rows = []
    for u in users:
        uid = u["telegram_id"]
        r, e = rdict.get(uid), edict.get(uid)
        status = "Topshirgan" if r else ("Ishlayapti" if e else "Topshirmagan")
        rows.append([u["full_name"], u["birth_date"], u["age"], u["grade"], u["school"], u["subject"], u["level"], status,
                     r["percent"] if r else None, r["correct"] if r and r["subject"] == MENTAL else None,
                     fmt_dur(r["duration_sec"]) if r else None, u["registered_at"], u["username"], uid])
    write_sheet(ws2, ["Ism-familya", "Tug‘ilgan sana", "Yosh", "Sinf", "Maktab", "Fan", "Daraja", "Holati",
                      "Foiz (%)", "Ball (Mental arifmetika)", "Sarflangan vaqt", "Ro‘yxatdan o‘tgan vaqt", "Username", "Telegram ID"], rows)

    # 3) Hozir ishlayotganlar
    ws3 = wb.create_sheet("Ishlayotganlar")
    rows = []
    for e in sorted(exams, key=lambda e: e["started_ts"]):
        u = udict.get(e["telegram_id"]) or {}
        rows.append([u.get("full_name"), e["subject"], e["level"], f"{e['idx']}/{len(json.loads(e['question_ids']))}",
                     stamp(e["started_ts"]), stamp(e["deadline_ts"]), e["telegram_id"]])
    write_sheet(ws3, ["Ism-familya", "Fan", "Daraja", "Javob berilgan savollar", "Boshlagan vaqt", "Tugash vaqti", "Telegram ID"], rows)

    # 4) Hali topshirmaganlar
    ws4 = wb.create_sheet("Topshirmaganlar")
    rows = [[u["full_name"], u["birth_date"], u["grade"], u["school"], u["subject"], u["level"], u["registered_at"], u["username"], u["telegram_id"]]
            for u in users if u["telegram_id"] not in rdict and u["telegram_id"] not in edict]
    write_sheet(ws4, ["Ism-familya", "Tug‘ilgan sana", "Sinf", "Maktab", "Fan", "Daraja", "Ro‘yxatdan o‘tgan vaqt", "Username", "Telegram ID"], rows)

    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    await update.message.reply_document(document=bio, filename=f"olimpiada_hisobot_{datetime.now(TZ).strftime('%Y%m%d_%H%M')}.xlsx",
        caption=f"Natijalar (foiz va sarflangan vaqt bilan), barcha o‘quvchilar, ishlayotganlar va topshirmaganlar.\n"
                f"Ro‘yxat: {len(users)} | Topshirgan: {len(results)} | Ishlayotgan: {len(exams)} | Topshirmagan: {len(users)-len(results)-len(exams)}")

async def text_router(update, context):
    text = (update.message.text or "").strip()
    if context.user_data.get("awaiting_schedule"):
        await schedule_save(update, context); return
    exam = get_exam(update.effective_user.id)
    if exam and exam["subject"] == MENTAL:
        await answer_typed(update, context, exam); return
    routes = {
        "📝 Ro‘yxatdan o‘tish": registration_start, "🔐 Admin panelga kirish": admin_begin,
        "🧪 Testni boshlash": test_start, "👤 Profil": profile,
        "📊 Statistika": admin_stats, "📥 Excel yuklab olish": export_excel,
        "✅ Topshirganlar": admin_finished, "🔄 Ishlayotganlar": admin_active, "⏳ Topshirmaganlar": admin_waiting,
        "🕒 Olimpiada vaqtini belgilash": schedule_prompt, "⏱ Davomiylikni o‘zgartirish": duration_prompt,
        "📅 Olimpiada vaqtini ko‘rish": schedule_show, "🚪 Admin paneldan chiqish": admin_exit,
    }
    if text in routes:
        await routes[text](update, context)
    elif text.startswith("🟢 Ro‘yxatdan o‘tishni"):
        await toggle_registration(update, context)
    else:
        await update.message.reply_text("Menyudagi tugmalardan foydalaning.",
                                        reply_markup=home_keyboard(update.effective_user.id) if get_user(update.effective_user.id) else main_keyboard())

async def error_handler(update, context):
    print("Xatolik:", repr(context.error))

def main():
    if not TOKEN:
        raise SystemExit("BOT_TOKEN ko‘rsatilmagan. README.md ni o‘qing.")
    if not DATABASE_URL:
        raise SystemExit("DATABASE_URL ko‘rsatilmagan (PostgreSQL manzili). README.md ni o‘qing.")
    if not ADMIN_PASSWORD or ADMIN_PASSWORD in ("change_this_password", "BU_YERGA_KUCHLI_PAROL_QOYING"):
        raise SystemExit("ADMIN_PASSWORD o‘rnatilmagan yoki standart qiymatda. Kuchli parol belgilang.")
    init_db()

    async def post_init(application):
        application.bot_data["watcher"] = asyncio.create_task(watcher(application))

    async def post_shutdown(application):
        POOL.close()

    app = Application.builder().token(TOKEN).post_init(post_init).post_shutdown(post_shutdown).build()
    reg_conv = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex("^📝 Ro‘yxatdan o‘tish$"), registration_start)],
        states={REGISTRATION: [MessageHandler(filters.TEXT & ~filters.COMMAND, registration_step)]},
        fallbacks=[CommandHandler("cancel", start)],
        allow_reentry=True)
    admin_conv = ConversationHandler(
        entry_points=[MessageHandler(filters.Regex("^🔐 Admin panelga kirish$"), admin_begin)],
        states={ADMIN_LOGIN_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_login)],
                ADMIN_PASSWORD_STATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, admin_password)]},
        fallbacks=[CommandHandler("cancel", start)],
        allow_reentry=True)
    app.add_handler(CommandHandler("start", start))
    app.add_handler(reg_conv)
    app.add_handler(admin_conv)
    app.add_handler(CallbackQueryHandler(answer_question, pattern=r"^answer:"))
    app.add_handler(CallbackQueryHandler(begin_callback, pattern=r"^begin:"))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, text_router))
    app.add_error_handler(error_handler)
    print("Olimpiada bot ishga tushdi.")
    app.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__ == "__main__":
    main()
