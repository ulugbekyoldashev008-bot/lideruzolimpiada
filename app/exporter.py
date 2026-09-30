from pathlib import Path
from tempfile import NamedTemporaryFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from .db import Answer, Attempt, Level, Participant, Question, Subject


async def create_excel(session) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Qatnashchilar"
    headers = ["№", "Qatnashchi ID", "F.I.O", "Telefon", "Viloyat", "Tuman", "Tug‘ilgan sana", "Yosh", "Fan", "Daraja", "Login", "Holati", "To‘g‘ri", "Noto‘g‘ri/javobsiz", "Ball", "Maksimal ball", "Foiz", "Admin izohi"]
    ws.append(headers)
    all_questions = (await session.execute(select(Question).where(Question.active.is_(True)))).scalars().all()
    questions_by_level = {}
    for question in all_questions:
        questions_by_level.setdefault((question.subject_id, question.level_id), []).append(question)
    all_answer_rows = (await session.execute(select(Answer, Question, Attempt, Participant).join(Question, Answer.question_id == Question.id).join(Attempt, Answer.attempt_id == Attempt.id).join(Participant, Attempt.participant_id == Participant.id).order_by(Participant.id, Question.position))).all()
    answers_by_attempt = {}
    for answer, question, attempt, participant in all_answer_rows:
        answers_by_attempt.setdefault(attempt.id, []).append((answer, question))
    rows = (await session.execute(select(Participant, Subject, Level, Attempt).join(Subject, Participant.subject_id == Subject.id).join(Level, Participant.level_id == Level.id).outerjoin(Attempt, Attempt.participant_id == Participant.id).order_by(Participant.id))).all()
    for i, (p, s, l, a) in enumerate(rows, 1):
        questions = questions_by_level.get((p.subject_id, p.level_id), [])
        answer_rows = answers_by_attempt.get(a.id, []) if a else []
        answer_by_question = {answer.question_id: answer for answer, _ in answer_rows}
        all_auto = bool(questions) and all(question.correct_option is not None or question.correct_text is not None for question in questions)
        correct = sum(1 for question in questions if answer_by_question.get(question.id) and (answer_by_question[question.id].score or 0) > 0) if all_auto else None
        incorrect = len(questions) - correct if correct is not None else None
        maximum = float(sum(question.max_score for question in questions)) if questions else None
        percent = (float(a.total_score or 0) / maximum * 100) if a and maximum else None
        ws.append([i, p.participant_code, p.full_name, p.phone, p.region, p.district, p.birth_date.strftime("%d.%m.%Y"), p.age, s.name, l.name, p.login, a.status if a else "ishlamagan", correct, incorrect, a.total_score if a else None, maximum, percent, a.admin_comment if a else None])
    ws2 = wb.create_sheet("Javoblar"); ws2.append(["Qatnashchi ID", "F.I.O", "Savol", "Javob", "Rasm file_id", "Ball", "Izoh"])
    for a, q, attempt, p in all_answer_rows: ws2.append([p.participant_code, p.full_name, q.text, a.text_answer, a.file_id, a.score, a.admin_comment])
    for sheet in (ws, ws2):
        for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1565C0"); cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            letter = col[0].column_letter; sheet.column_dimensions[letter].width = min(45, max(12, max(len(str(c.value or "")) for c in col) + 2))
    temp = NamedTemporaryFile(prefix="olimpiada_", suffix=".xlsx", delete=False); temp.close(); wb.save(temp.name)
    return Path(temp.name)
